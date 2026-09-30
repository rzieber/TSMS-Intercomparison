import numpy as np
import pandas as pd
from pathlib import Path
from functools import reduce
from scipy.stats import pearsonr
from sklearn.metrics import mean_absolute_error, root_mean_squared_error


output = Path("data/error-analysis")
data = Path("data/cleaned")

TIMESCALE = 'h'    # 'h' --> hourly, 'D' --> daily, None --> point-for-point
# Command line override (main.py uses it): python scripts/error/error-analysis.py --timescale D   (h | D | none)
import argparse
_cli = argparse.ArgumentParser(description="3D-PAWS vs TSMS error analysis and WMO classification")
_cli.add_argument("--timescale", choices=["h", "D", "none"], help="aggregation before statistics (default: h)")
_timescale = _cli.parse_args().timescale
if _timescale is not None:
    TIMESCALE = None if _timescale == "none" else _timescale

# Wind regimes, classified from the TSMS reference only (never from the 3D-PAWS instrument under test).
# Non-variable: reference 10-m speed >= 3.0 m/s (~6 kt), the same cutoff used for the wind-rose plots.
# Variable: reference 10-m speed < 3.0 m/s. Each wind variable is analysed for all winds and for each regime.
NONVARIABLE_THRESHOLD = 3.0     # m/s, applied to the reference's unadjusted 10-m avg_wind_speed
wind_regimes = ["variable", "nonvariable"]

variable_mapper = { # TSMS : 3DPAWS
    "temperature":                  ["bmp2_temp", "htu_temp", "sth_temp", "mcp9808"],
    "humidity":                     ["bme2_hum", "htu_hum", "sth_hum"],
    "actual_pressure":              ["bmp2_pres"],
    "avg_wind_dir":                 ["wind_dir"],
    "avg_wind_speed_2m":            ["wind_speed"],
    "avg_wind_dir_variable":        ["wind_dir_variable"],
    "avg_wind_speed_2m_variable":   ["wind_speed_variable"],
    "avg_wind_dir_nonvariable":     ["wind_dir_nonvariable"],
    "avg_wind_speed_2m_nonvariable":["wind_speed_nonvariable"],
    "total_rainfall":               ["tipping"]
}
# WMO criteria (Annex 1.A / 1.G, OSCAR) are defined in the WMO classification section at the end of this script.
magnetic_declinations = { # https://www.ngdc.noaa.gov/geomag/calculators/magcalc.shtml
    "Ankara":   6.11,          
    "Konya":    5.73,            
    "Adana":    5.69            
}
sht_upgrade = { 
    "Ankara":   pd.to_datetime("2024-01-17 00:00:00"), 
    "Konya":    pd.to_datetime("2024-01-12 00:00:00"),  
    "Adana":    pd.to_datetime("2024-01-15 00:00:00")   
}
hellman_exponents = {
    "Ankara": 0.30,  # Urban
    "Konya":  0.35,  # Urban + many obstacles  
    "Adana":  0.25   # Suburban + grass
}
# height in meters of wind sensors
h_tsms = 10.0 
h_paws = 2.0   


"""
=============================================================================================================================
Create the 3D PAWS and TSMS dataframes.
=============================================================================================================================    
"""
paws_dfs = {
    "Ankara":[],
    "Konya" :[],
    "Adana" :[]
}
tsms_dfs = {}

for csv in sorted(data.glob("*_final.csv")):   # top level only, so backup subfolders are never read
    name = str(csv.stem)
    df = pd.read_csv(csv)

    key = "_".join(name.split("_")[2:3])

    if name.startswith("3DPAWS"):
        station_name = "_".join(name.split("_")[1:2])
        print("Reading", station_name)
        paws_dfs[key].append((df, station_name))
    elif name.startswith("TSMS"):
        print("Reading", key, "reference")
        # Hellman correction
        df['avg_wind_speed_2m'] = df['avg_wind_speed'] * \
            ((h_paws / h_tsms) ** hellman_exponents[key])
        tsms_dfs[key] = df

print()


"""
=============================================================================================================================
Wind: remove meaningless directions and split into variable / non-variable regimes (minute level, before aggregation).
    - A direction recorded with a speed of 0 carries no information: 3D-PAWS keeps reporting the vane's last position,
      TSMS reports 0. Those directions are set to NaN. Zero SPEEDS are kept -- a 3D-PAWS cup reading 0 while the
      reference shows wind is a genuine measurement error, not a calm.
    - Each minute is assigned a regime from the reference's 10-m speed at the same timestamp. 3D-PAWS minutes with no
      reference minute get no regime (they still count toward the all-winds statistics).
=============================================================================================================================
"""
base_wind_pairs = {  # direction column : speed column
    'avg_wind_dir': 'avg_wind_speed_2m',    # TSMS
    'wind_dir':     'wind_speed'            # 3D-PAWS
}
wind_pairs = dict(base_wind_pairs)
for regime in wind_regimes:
    for dir_col, spd_col in base_wind_pairs.items():
        wind_pairs[f'{dir_col}_{regime}'] = f'{spd_col}_{regime}'

def _mask_stale_direction(df:pd.DataFrame):
    """Set wind direction to NaN wherever the same instrument's wind speed is 0 (or missing)."""
    df = df.copy()
    for dir_col, spd_col in base_wind_pairs.items():
        if dir_col in df.columns and spd_col in df.columns:
            df.loc[~(df[spd_col] > 0), dir_col] = np.nan
    return df

def _add_wind_regimes(df:pd.DataFrame, ref_speed:pd.Series):
    """Add <wind column>_variable / _nonvariable copies, blanked outside each regime. ref_speed: reference 10-m speed by date."""
    df = df.copy()
    df['date'] = pd.to_datetime(df['date'])
    ref = df['date'].map(ref_speed)
    regime_mask = {'variable': ref < NONVARIABLE_THRESHOLD, 'nonvariable': ref >= NONVARIABLE_THRESHOLD}
    for regime in wind_regimes:
        for dir_col, spd_col in base_wind_pairs.items():
            if dir_col in df.columns and spd_col in df.columns:
                df[f'{dir_col}_{regime}'] = df[dir_col].where(regime_mask[regime])
                df[f'{spd_col}_{regime}'] = df[spd_col].where(regime_mask[regime])
    return df

for key in tsms_dfs:
    tsms_dfs[key]['date'] = pd.to_datetime(tsms_dfs[key]['date'])
    ref_speed = tsms_dfs[key].set_index('date')['avg_wind_speed']

    tsms_dfs[key] = _add_wind_regimes(_mask_stale_direction(tsms_dfs[key]), ref_speed)
    paws_dfs[key] = [(_add_wind_regimes(_mask_stale_direction(df), ref_speed), name) for df, name in paws_dfs[key]]

    n_ref = ref_speed.notna().sum()
    print(f"{key} reference: {100 * (ref_speed >= NONVARIABLE_THRESHOLD).sum() / n_ref:.1f}% of minutes non-variable "
          f"(>= {NONVARIABLE_THRESHOLD} m/s at 10 m)")

print()


"""
=============================================================================================================================
Pair each 3D-PAWS variable with its reference variable at MINUTE level, then aggregate the pairs.
    Only minutes where both instruments report are used, so both sides of every hourly/daily value cover exactly the
    same minutes. An hour/day is kept only if at least PAIRED_COMPLETENESS of its minutes are paired. For wind variables
    this is judged on minutes where both instruments report a wind speed, so the regime split and the removal of
    zero-speed directions don't make windy/calm periods look incomplete. (Each instrument
    has already passed the per-instrument 80% daily completeness check in outlier-removal.py Phase 7; two instruments
    can each pass it and still share far fewer minutes.)
    Aggregation: mean (temperature, humidity, pressure, wind speed), sum (rainfall), speed-weighted vector mean
    (wind direction, each instrument weighted by its own speed).
=============================================================================================================================
"""
PAIRED_COMPLETENESS = 0.8
minutes_per_period = {'10MIN': 10, 'H': 60, 'D': 1440}
sum_variables = {'total_rainfall', 'tipping', 'total_rainfall_nf', 'tipping_nf'}
_pair_cache = {}
_wind_availability_cache = {}

def _wind_availability(station:str, tsms_df:pd.DataFrame, paws_df:pd.DataFrame, freq):
    """Paired minutes per period in which BOTH instruments report a wind speed (zeros count as reporting).
    Completeness for every wind variable is judged on this, not on the regime- or calm-filtered values: a
    non-variable-winds hour with 10 windy minutes is a complete hour if both anemometers reported all hour."""
    if (station, freq) not in _wind_availability_cache:
        both = pd.merge(tsms_df[['date', 'avg_wind_speed_2m']], paws_df[['date', 'wind_speed']], on='date', how='inner').dropna()
        _wind_availability_cache[(station, freq)] = both.set_index('date')['wind_speed'].resample(freq).count()
    return _wind_availability_cache[(station, freq)]

def _paired(station:str, tsms_df:pd.DataFrame, paws_df:pd.DataFrame, variable:str, var:str, freq='timescale'):
    """Return the paired (and, if freq is set, aggregated) values as a DataFrame with columns date, tsms, paws.
    freq defaults to TIMESCALE; the WMO classification passes its own averaging times (None = 1-min pairs)."""
    if freq == 'timescale': freq = TIMESCALE
    cache_key = (station, var, freq)
    if cache_key in _pair_cache: return _pair_cache[cache_key]

    is_direction = var in wind_pairs
    t = tsms_df[['date', variable]].rename(columns={variable: 'tsms'})
    p = paws_df[['date', var]].rename(columns={var: 'paws'})
    if is_direction:
        t['tsms_spd'] = tsms_df[wind_pairs[variable]]
        p['paws_spd'] = paws_df[wind_pairs[var]]
    merged = pd.merge(t, p, on='date', how='inner').dropna(subset=['tsms', 'paws'])

    if freq is not None:
        g = merged.set_index('date')
        n_paired = g['tsms'].resample(freq).count()

        if is_direction:    # speed-weighted vector mean of the paired minutes
            comps = pd.DataFrame({
                'tu': g['tsms_spd'] * np.sin(np.radians(g['tsms'])), 'tv': g['tsms_spd'] * np.cos(np.radians(g['tsms'])),
                'pu': g['paws_spd'] * np.sin(np.radians(g['paws'])), 'pv': g['paws_spd'] * np.cos(np.radians(g['paws']))
            }).resample(freq).mean()
            agg = pd.DataFrame({
                'tsms': np.degrees(np.arctan2(comps['tu'], comps['tv'])).round(6) % 360,
                'paws': np.degrees(np.arctan2(comps['pu'], comps['pv'])).round(6) % 360
            })
        elif var in sum_variables:
            agg = g[['tsms', 'paws']].resample(freq).sum()
        else:
            agg = g[['tsms', 'paws']].resample(freq).mean()

        if var.startswith('wind_'):     # availability of wind data, not of the regime/calm-filtered values
            n_paired = _wind_availability(station, tsms_df, paws_df, freq).reindex(agg.index).fillna(0)
        complete = n_paired >= PAIRED_COMPLETENESS * minutes_per_period[str(freq).upper()]
        merged = agg[complete].dropna().reset_index()

    _pair_cache[cache_key] = merged[['date', 'tsms', 'paws']]
    return _pair_cache[cache_key]

print(f"Pairing at minute level" + (f", then aggregating to {TIMESCALE} (>= {PAIRED_COMPLETENESS:.0%} paired minutes)" if TIMESCALE else ""))

"""
=============================================================================================================================
Difference Analysis ----------------------------------
    Mean Bias, Median Bias, Standard Deviation and Interquartile Range of the differences (3D-PAWS - TSMS)
=============================================================================================================================
"""
difference = []  # [TSMS00[TSMS00, var, mean, median, std, iqr], TSMS01[TSMS01, var, mean, median, std, iqr], ...]
difference.append(['Station', 'Variable', 'Mean Bias', 'Median Bias', 'Bias Standard Deviation', 'Bias Interquartile Range'])

print("Difference statistics.")
for key in paws_dfs:
    tsms_df = tsms_dfs[key]

    for paws_df, stn_name in paws_dfs[key]:
        print(stn_name)

        for variable in variable_mapper.keys():
            print(f'\t\t{variable}')

            for var in variable_mapper[variable]:
                merged = None
                if var in ['bme2_hum', 'bmp2_slp']: continue
                merged = _paired(stn_name, tsms_df, paws_df, variable, var)

                tsms = merged['tsms'].to_numpy()
                paws = merged['paws'].to_numpy()

                if len(paws) == 0:
                    difference.append([stn_name, var, np.nan, np.nan, np.nan, np.nan])
                    continue

                if var.startswith('wind_dir'): # signed circular difference in [-180, 180)
                    diff = ((paws - tsms) + 180) % 360 - 180
                elif var in ['tipping']: # exclude dry periods
                    nonzero_mask = ~((tsms == 0) & (paws == 0))
                    diff = paws[nonzero_mask] - tsms[nonzero_mask]
                else:
                    diff = paws - tsms

                q1, q3 = np.percentile(diff, [25, 75])
                difference.append([
                    stn_name, var,
                    round(diff.mean(), 1),                  # mean bias
                    round(np.median(diff), 1),              # median bias
                    round(np.std(diff, ddof=1), 1),         # standard deviation
                    round(q3 - q1, 1)                       # interquartile range
                ])

df_difference = pd.DataFrame(difference[1:], columns=difference[0])

print()

"""
=============================================================================================================================
Accuracy (systematic errors) -------------------------
    Mean Absolute Error, Root Mean Squared Error
=============================================================================================================================
"""
accuracy = []  # [TSMS00[TSMS00, var, mae, rmse], TSMS01[TSMS01, var, mae, rmse], ...]
accuracy.append(['Station', 'Variable', 'Mean Absolute Error', 'Root Mean Squared Error'])

print("Accuracy statistics.")
for key in paws_dfs:
    tsms_df = tsms_dfs[key]

    for df in paws_dfs[key]:
        print(df[1])
        paws_df = df[0]

        for variable in variable_mapper.keys():
            print(f'\t\t{variable}')

            for var in variable_mapper[variable]:
                merged = None
                if var in ['bme2_hum', 'bmp2_slp']: continue
                merged = _paired(df[1], tsms_df, paws_df, variable, var)

                tsms = merged['tsms'].to_numpy()
                paws = merged['paws'].to_numpy()

                if len(paws) == 0: 
                    accuracy.append([df[1], f'{var}', np.nan, np.nan])
                    continue     

                s = []
                s.append(df[1])
                s.append(f'{var}')

                if var.startswith('wind_dir'): # handle circular discrepancy at 360 -> 0
                    # paws_corrected = (paws + magnetic_declinations[key] + wind_dir_bias[key]) % 360
                    # raw_diff = paws_corrected - tsms
                    raw_diff = paws - tsms
                    abs_err = np.minimum(np.abs(raw_diff), 360 - np.abs(raw_diff))
                    s.append(round(abs_err.mean(), 1))                  # mae
                    s.append(round(np.sqrt((abs_err**2).mean()), 1))    # rmse
                elif var in ['tipping']: # exclude dry periods
                    nonzero_mask = ~((tsms == 0) & (paws == 0))
                    tsms_nonzero = tsms[nonzero_mask]
                    paws_nonzero = paws[nonzero_mask]
                    s.append(round(mean_absolute_error(tsms_nonzero, paws_nonzero), 1))
                    s.append(round(root_mean_squared_error(tsms_nonzero, paws_nonzero),1))
                else:
                    s.append(round(mean_absolute_error(tsms, paws), 1))
                    s.append(round(root_mean_squared_error(tsms, paws),1))     

                accuracy.append(s)

df_accuracy = pd.DataFrame(accuracy[1:], columns=accuracy[0])
# df_accuracy.to_csv(output / "accuracy.csv", index=False)

print()

"""
=============================================================================================================================
Precision (random errors) ----------------------
    Standard Deviations, Correlation Coefficients
=============================================================================================================================    
"""
precision = [] # [TSMS00[TSMS00, var, std_paws, std_tsms, r], TSMS01[TSMS01, var, std, r], ...]
precision.append(['Station', 'Variable', '3D-PAWS Standard Deviation', 'TSMS Standard Deviation', 'Pearson Correlation Coefficient'])

print("Precision statistics.")
for key in paws_dfs:
    tsms_df = tsms_dfs[key]

    for paws_df, stn_name in paws_dfs[key]:
        print(stn_name)

        for variable in variable_mapper.keys():
            print(f'\t\t{variable}')

            for var in variable_mapper[variable]:
                merged = None
                if var in ['bme2_hum', 'bmp2_slp']: continue
                merged = _paired(stn_name, tsms_df, paws_df, variable, var)

                tsms = merged['tsms'].to_numpy()
                paws = merged['paws'].to_numpy()

                if len(paws) == 0: 
                    precision.append([stn_name, var, np.nan, np.nan, np.nan])
                    continue    

                if var in ['tipping']:
                    nonzero_mask = ~((tsms == 0) & (paws == 0))
                    tsms_nonzero = tsms[nonzero_mask]
                    paws_nonzero = paws[nonzero_mask]
                    corr_coef, _ = pearsonr(tsms_nonzero, paws_nonzero)
                    paws_std = np.std(paws_nonzero, ddof=1)
                    tsms_std = np.std(tsms_nonzero, ddof=1)
                elif var.startswith('wind_dir'):
                    # paws_corrected = (paws + magnetic_declinations[key] + wind_dir_bias[key]) % 360
                    # corr_coef, _ = pearsonr(tsms, paws_corrected)
                    corr_coef, _ = pearsonr(tsms, paws)
                    # paws_std = np.std(paws_corrected, ddof=1)
                    paws_std = np.std(paws, ddof=1)
                    tsms_std = np.std(tsms, ddof=1)
                else:
                    corr_coef, _ = pearsonr(tsms, paws)
                    paws_std = np.std(paws, ddof=1)
                    tsms_std = np.std(tsms, ddof=1)

                precision.append([stn_name, var, round(paws_std, 1), round(tsms_std, 1), round(corr_coef, 1)])

df_precision = pd.DataFrame(precision[1:], columns=precision[0])
# df_precision.to_csv(output / "precision.csv", index=False)

print()

"""
=============================================================================================================================
Create comprehensive csv
=============================================================================================================================    
"""
print("Creating CSV.")
key_cols = ['Station', 'Variable']

# For each DataFrame, select only key columns + non-key columns (excluding duplicates)
def select_merge_cols(df):
    return df[key_cols + [col for col in df.columns if col not in key_cols]]

dfs = [df_difference, df_accuracy, df_precision]
dfs_selected = [select_merge_cols(df) for df in dfs]

df_merged = reduce(
    lambda left, right: pd.merge(left, right, on=key_cols, how='outer'),
    dfs_selected
)

# order stations numerically (TSMS00, TSMS01, ...) and variables as listed in variable_mapper
station_order = sorted(df_merged['Station'].unique(), key=lambda s: int(''.join(filter(str.isdigit, s))))
variable_order = [v for vs in variable_mapper.values() for v in vs if v in df_merged['Variable'].unique()]
df_merged['Station'] = pd.Categorical(df_merged['Station'], categories=station_order, ordered=True)
df_merged['Variable'] = pd.Categorical(df_merged['Variable'], categories=variable_order, ordered=True)
df_merged = df_merged.sort_values(by=['Station', 'Variable']).reset_index(drop=True)

if      str(TIMESCALE).upper() == 'D':      filename = 'daily-statistical-analysis'
elif    str(TIMESCALE).upper() == 'H':      filename = 'hourly-statistical-analysis'
else:                                       filename = 'point-for-point-statistical-analysis'

df_merged.to_csv(output / f'{filename}.csv', index=False)


"""
=============================================================================================================================
Create Excel workbook: one sheet per statistic (rows = variable, columns = station), plus the full table
=============================================================================================================================
"""
print("Creating Excel workbook.")
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

stat_cols = [col for col in df_merged.columns if col not in key_cols]
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="305496")

def _format_sheet(ws, first_col_width=18, col_width=12):
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"
    ws.column_dimensions['A'].width = first_col_width
    for i in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = col_width

with pd.ExcelWriter(output / f'{filename}.xlsx', engine='openpyxl') as writer:
    for stat in stat_cols:
        pivot = df_merged.pivot(index='Variable', columns='Station', values=stat)
        pivot.index.name = 'Variable'
        pivot.columns.name = None
        pivot.to_excel(writer, sheet_name=stat[:31])

        ws = writer.sheets[stat[:31]]
        _format_sheet(ws)
        ws.row_dimensions[1].height = 20

    df_merged.to_excel(writer, sheet_name='All Stations', index=False)
    ws = writer.sheets['All Stations']
    _format_sheet(ws, first_col_width=10, col_width=14)
    ws.freeze_panes = "C2"
    ws.row_dimensions[1].height = 45
    ws.auto_filter.ref = ws.dimensions


"""
=============================================================================================================================
WMO classification (docs/potential-fixes.md PF-15) -- independent of TIMESCALE
    Tier 1: WMO-No. 8 Vol. I, Ch. 1, Annex 1.G Measurement Quality Classification (Class A/B/C, else D), with the
            Annex 1.A required / achievable uncertainty shown for context. All at 95% (k = 2): a class is met when
            >= 95% of paired values have |3D-PAWS - TSMS| within the class's target system uncertainty.
    Tier 2: OSCAR/Requirements for selected application areas (docs/oscar-requirements.md). OSCAR uncertainty is an
            RMSE at 68% (k = 1), so the 3D-PAWS RMSE is compared with goal / breakthrough / threshold.
    Values are compared at the WMO averaging time (Annex 1.A): 1-min values for temperature, humidity and pressure,
    10-min averages for wind, daily totals for rain (OSCAR NWP/hydrology rain: hourly totals).
    Conservative by construction: the differences include the TSMS reference's own uncertainty and siting effects,
    which the WMO scheme treats separately (Annex 1.D). Konya and Ankara wind is flagged as siting-limited.
=============================================================================================================================
"""
print("WMO classification (Annex 1.G / Annex 1.A / OSCAR).")

def _const(x): return lambda r: np.full(len(r), float(x))

def _measurand(variable:str):
    if variable in ('temperature', 'humidity'): return variable
    if variable == 'actual_pressure':           return 'pressure'
    if variable.startswith('avg_wind_speed'):   return 'wind_speed'
    if variable.startswith('avg_wind_dir'):     return 'wind_dir'
    if variable == 'total_rainfall':            return 'precip'
    return None

WMO_AVERAGING = {'temperature': None, 'humidity': None, 'pressure': None, 'wind_speed': '10min', 'wind_dir': '10min', 'precip': 'D'}
AVERAGING_LABEL = {None: '1 min', '10min': '10-min mean', 'h': 'hourly total', 'D': 'daily total'}
WET_THRESHOLD = {'D': 0.2, 'h': 0.0}    # rain periods kept when either gauge exceeds this (dry-dry periods excluded)

ANNEX_1G = {    # Vol. I, Ch. 1, Annex 1.G, Table 1: target system uncertainty (k = 2) for Class A, B, C, as f(reference value)
    'temperature':  [_const(0.2), _const(0.6), _const(1.0)],
    'humidity':     [_const(2), _const(5), _const(10)],
    'pressure':     [_const(0.2), _const(1.0), _const(2.0)],
    'wind_speed':   [lambda r: np.maximum(1, 0.05 * r), lambda r: np.maximum(2, 0.10 * r), lambda r: np.maximum(5, 0.15 * r)],
    'wind_dir':     [_const(5), _const(10), _const(15)],
    'precip':       [lambda r: np.maximum(1, 0.02 * r), lambda r: np.maximum(3, 0.05 * r), lambda r: np.maximum(5, 0.10 * r)],
}
ANNEX_1A_REQUIRED = {   # Vol. I, Ch. 1, Annex 1.A, column 5 (k = 2)
    'temperature':  lambda r: np.where((r > -40) & (r <= 40), 0.1, 0.3),
    'humidity':     _const(1),
    'pressure':     _const(0.1),
    'wind_speed':   lambda r: np.where(r <= 5, 0.5, 0.10 * r),
    'wind_dir':     _const(5),
    'precip':       lambda r: np.where(r <= 5, 0.1, 0.02 * r),
}
ANNEX_1A_ACHIEVABLE = { # Vol. I, Ch. 1, Annex 1.A, column 8 (k = 2); wind speed has no numeric value there
    'temperature':  _const(0.2),
    'humidity':     _const(3),      # solid-state sensors
    'pressure':     _const(0.15),
    'wind_dir':     _const(5),
    'precip':       lambda r: np.maximum(0.1, 0.05 * r),
}
OSCAR = {       # OSCAR/Requirements (k = 1 RMSE): (goal, breakthrough, threshold, averaging), docs/oscar-requirements.md
    '2.2 High-Res NWP':     {'temperature': (0.5, 1, 3, None), 'wind_speed': (0.5, 1, 3, '10min'), 'precip': (0.5, 2, 5, 'h')},
    '2.3 Nowcasting/VSRF':  {'temperature': (0.5, 1, 3, None), 'wind_speed': (1, 1.4, 3, '10min')},
    '2.9 Agricultural Met': {'precip': (2, 5, 10, 'D')},
    '4.1 Hydrology':        {'precip': (0.5, 2, 5, 'h')},
}
SITING_LIMITED_WIND = {  # docs/siting.md -- the class is NOT a measure of 3D-PAWS sensor performance at these sites
    'Konya':  "Siting-limited (courtyard next to a brick wall, docs/siting.md). Not a measure of sensor performance: "
              "sheltered, weak 2-m winds make absolute speed errors small (high classes are an artifact) and direction unreliable",
    'Ankara': "Siting-limited (nearby hill affects 2-m wind, docs/siting.md). Class reflects siting as well as the sensor; "
              "don't attribute it to sensor performance alone",
}

def _differences(station, tsms_df, paws_df, variable, var, measurand, freq):
    """Paired differences (3D-PAWS - TSMS) and reference values at the given averaging time."""
    m = _paired(station, tsms_df, paws_df, variable, var, freq=freq)
    t, p = m['tsms'].to_numpy(), m['paws'].to_numpy()
    if measurand == 'precip':
        wet = np.maximum(t, p) > WET_THRESHOLD[freq]
        t, p = t[wet], p[wet]
    d = ((p - t) + 180) % 360 - 180 if measurand == 'wind_dir' else p - t
    return d, t

wmo_rows = []
for key in paws_dfs:
    tsms_df = tsms_dfs[key]
    for paws_df, stn_name in paws_dfs[key]:
        for variable in variable_mapper.keys():
            measurand = _measurand(variable)
            if measurand is None: continue
            for var in variable_mapper[variable]:
                if var in ['bme2_hum', 'bmp2_slp']: continue
                freq = WMO_AVERAGING[measurand]
                d, ref = _differences(stn_name, tsms_df, paws_df, variable, var, measurand, freq)
                row = {'Station': stn_name, 'Site': key, 'Variable': var, 'Measurand': measurand,
                       'Averaging': AVERAGING_LABEL[freq], 'N': len(d)}
                if len(d) == 0:
                    wmo_rows.append(row); continue
                absd = np.abs(d)
                row.update({'Bias': round(d.mean(), 2), 'RMSE (k=1)': round(np.sqrt((d ** 2).mean()), 2),
                            'U95 (95th pct |diff|)': round(np.percentile(absd, 95), 2)})
                achieved = 'D'
                for cls, tol in zip('ABC', ANNEX_1G[measurand]):
                    within = (absd <= tol(ref)).mean()
                    row[f'% within Class {cls}'] = round(100 * within, 1)
                    if achieved == 'D' and within >= 0.95: achieved = cls
                row['Annex 1.G class'] = achieved
                row['% within Annex 1.A required'] = round(100 * (absd <= ANNEX_1A_REQUIRED[measurand](ref)).mean(), 1)
                if measurand in ANNEX_1A_ACHIEVABLE:
                    row['% within Annex 1.A achievable'] = round(100 * (absd <= ANNEX_1A_ACHIEVABLE[measurand](ref)).mean(), 1)
                for area, reqs in OSCAR.items():
                    if measurand not in reqs: continue
                    goal, breakthrough, threshold, area_freq = reqs[measurand]
                    d_area, _ = _differences(stn_name, tsms_df, paws_df, variable, var, measurand, area_freq)
                    if len(d_area) == 0: continue
                    rmse = np.sqrt((d_area ** 2).mean())
                    level = 'goal' if rmse <= goal else 'breakthrough' if rmse <= breakthrough else \
                            'threshold' if rmse <= threshold else 'below threshold'
                    row[f'OSCAR {area}'] = f"{level} (RMSE {rmse:.2f}, {AVERAGING_LABEL[area_freq]})"
                row['Siting note'] = SITING_LIMITED_WIND.get(key, '') if measurand.startswith('wind') else ''
                wmo_rows.append(row)

        # Rain again with flagged days excluded: a day is dropped when either gauge has any QC flag that day
        # (tipping_flag / total_rainfall_flag, docs/qc-framework.md Step 8), e.g. SF-09 Ankara reference, SF-12 TSMS05.
        day_t = pd.to_datetime(tsms_df['date']).dt.floor('D'); day_p = pd.to_datetime(paws_df['date']).dt.floor('D')
        flagged = set(day_t[tsms_df['total_rainfall_flag'].notna().to_numpy()]) | set(day_p[paws_df['tipping_flag'].notna().to_numpy()])
        t_nf = tsms_df.assign(total_rainfall_nf=tsms_df['total_rainfall'].where(~day_t.isin(flagged)))
        p_nf = paws_df.assign(tipping_nf=paws_df['tipping'].where(~day_p.isin(flagged)))
        freq = WMO_AVERAGING['precip']
        d, ref = _differences(stn_name, t_nf, p_nf, 'total_rainfall_nf', 'tipping_nf', 'precip', freq)
        row = {'Station': stn_name, 'Site': key, 'Variable': 'tipping (flagged days excluded)', 'Measurand': 'precip',
               'Averaging': AVERAGING_LABEL[freq], 'N': len(d)}
        if len(d):
            absd = np.abs(d)
            row.update({'Bias': round(d.mean(), 2), 'RMSE (k=1)': round(np.sqrt((d ** 2).mean()), 2),
                        'U95 (95th pct |diff|)': round(np.percentile(absd, 95), 2)})
            achieved = 'D'
            for cls, tol in zip('ABC', ANNEX_1G['precip']):
                within = (absd <= tol(ref)).mean()
                row[f'% within Class {cls}'] = round(100 * within, 1)
                if achieved == 'D' and within >= 0.95: achieved = cls
            row['Annex 1.G class'] = achieved
            row['% within Annex 1.A required'] = round(100 * (absd <= ANNEX_1A_REQUIRED['precip'](ref)).mean(), 1)
            row['% within Annex 1.A achievable'] = round(100 * (absd <= ANNEX_1A_ACHIEVABLE['precip'](ref)).mean(), 1)
        wmo_rows.append(row)

df_wmo = pd.DataFrame(wmo_rows)
df_wmo['Station'] = pd.Categorical(df_wmo['Station'], categories=station_order, ordered=True)
df_wmo['Variable'] = pd.Categorical(df_wmo['Variable'], categories=variable_order + ['tipping (flagged days excluded)'], ordered=True)
df_wmo = df_wmo.sort_values(['Station', 'Variable']).reset_index(drop=True)
df_wmo.to_csv(output / 'wmo-classification.csv', index=False)

criteria = pd.DataFrame([
    {'Measurand': 'temperature', 'Averaging': '1 min', 'Annex 1.G A / B / C (k=2)': '0.2 / 0.6 / 1.0 K', 'Annex 1.A required (k=2)': '0.1 K (-40..40 °C), else 0.3 K', 'Annex 1.A achievable (k=2)': '0.2 K', 'OSCAR goal / breakthrough / threshold (k=1)': '2.2 & 2.3: 0.5 / 1 / 3 K (PBL layer)'},
    {'Measurand': 'humidity', 'Averaging': '1 min', 'Annex 1.G A / B / C (k=2)': '2 / 5 / 10 %RH', 'Annex 1.A required (k=2)': '1 %RH', 'Annex 1.A achievable (k=2)': '3 %RH (solid state)', 'OSCAR goal / breakthrough / threshold (k=1)': 'not compared (OSCAR uses specific humidity)'},
    {'Measurand': 'pressure', 'Averaging': '1 min', 'Annex 1.G A / B / C (k=2)': '0.2 / 1.0 / 2.0 hPa', 'Annex 1.A required (k=2)': '0.1 hPa', 'Annex 1.A achievable (k=2)': '0.15 hPa', 'OSCAR goal / breakthrough / threshold (k=1)': 'no near-surface pressure requirement found'},
    {'Measurand': 'wind_speed', 'Averaging': '10 min', 'Annex 1.G A / B / C (k=2)': 'max(1 m/s, 5%) / max(2 m/s, 10%) / max(5 m/s, 15%)', 'Annex 1.A required (k=2)': '0.5 m/s (<= 5 m/s), 10% above', 'Annex 1.A achievable (k=2)': 'not given', 'OSCAR goal / breakthrough / threshold (k=1)': '2.2: 0.5 / 1 / 3 m/s; 2.3: 1 / 1.4 / 3 m/s'},
    {'Measurand': 'wind_dir', 'Averaging': '10 min', 'Annex 1.G A / B / C (k=2)': '5 / 10 / 15 deg', 'Annex 1.A required (k=2)': '5 deg', 'Annex 1.A achievable (k=2)': '5 deg', 'OSCAR goal / breakthrough / threshold (k=1)': 'not compared'},
    {'Measurand': 'precip', 'Averaging': 'daily total (wet days)', 'Annex 1.G A / B / C (k=2)': 'max(1 mm, 2%) / max(3 mm, 5%) / max(5 mm, 10%)', 'Annex 1.A required (k=2)': '0.1 mm (<= 5 mm), 2% above', 'Annex 1.A achievable (k=2)': 'max(0.1 mm, 5%)', 'OSCAR goal / breakthrough / threshold (k=1)': '2.9 (daily): 2 / 5 / 10 mm; 2.2 & 4.1 (hourly): 0.5 / 2 / 5 mm'},
])
with pd.ExcelWriter(output / 'wmo-classification.xlsx', engine='openpyxl') as writer:
    df_wmo.to_excel(writer, sheet_name='Classification', index=False)
    _format_sheet(writer.sheets['Classification'], first_col_width=10, col_width=16)
    writer.sheets['Classification'].freeze_panes = "D2"
    writer.sheets['Classification'].row_dimensions[1].height = 45
    writer.sheets['Classification'].auto_filter.ref = writer.sheets['Classification'].dimensions
    criteria.to_excel(writer, sheet_name='Criteria', index=False)
    _format_sheet(writer.sheets['Criteria'], first_col_width=14, col_width=34)

print(f"Wrote wmo-classification.csv / .xlsx ({len(df_wmo)} rows)")
